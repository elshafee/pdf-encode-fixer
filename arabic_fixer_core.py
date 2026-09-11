#!/usr/bin/env python3
"""
arabic_fixer_core.py
High-precision engine to extract, restore, and generate academic Mark List PDFs and Excel files
with reconstructed Arabic text and executive layout.
"""

import os
import re
import io
import fitz
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

import reportlab
from reportlab.lib.pagesizes import A4, landscape
from reportlab.platypus import (
    SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle, KeepTogether, HRFlowable
)
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.lib import colors
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
import arabic_reshaper
from bidi.algorithm import get_display

# ---------------------------------------------------------------------------
# Font Registration
# ---------------------------------------------------------------------------
def setup_arabic_font():
    font_candidates = [
        ('/System/Library/Fonts/Supplemental/Arial Unicode.ttf', 'ArialUnicode'),
        ('/Library/Fonts/Arial Unicode.ttf', 'ArialUnicode'),
        ('/System/Library/Fonts/Supplemental/Tahoma.ttf', 'Tahoma'),
        ('/System/Library/Fonts/Supplemental/Arial.ttf', 'Arial'),
    ]
    registered_name = None
    for path, name in font_candidates:
        if os.path.exists(path):
            try:
                pdfmetrics.registerFont(TTFont('ArabicAcademicFont', path))
                registered_name = 'ArabicAcademicFont'
                break
            except Exception:
                continue
    if not registered_name:
        registered_name = 'Helvetica'
    return registered_name

ARABIC_FONT_NAME = setup_arabic_font()

def ar(text):
    """Reshape and reorder Arabic text for ReportLab PDF rendering."""
    if text is None:
        return ''
    s = str(text).strip()
    if not s:
        return ''
    # Check if text contains Arabic characters
    if any(0x0600 <= ord(c) <= 0x06FF or 0xFB50 <= ord(c) <= 0xFDFF or 0xFE70 <= ord(c) <= 0xFEFF for c in s):
        try:
            reshaped = arabic_reshaper.reshape(s)
            return get_display(reshaped)
        except Exception:
            return s
    return s

def format_grade(g):
    """Formats bilingual grades cleanly, e.g. 'ج|C' -> 'C (ج)'."""
    if not g:
        return '-'
    g = str(g).strip()
    if '|' in g:
        parts = g.split('|')
        ar_part = ar(parts[0].strip())
        en_part = parts[1].strip()
        return f"{en_part} ({ar_part})"
    return ar(g)

# ---------------------------------------------------------------------------
# Mojibake & Arabic Reconstruction
# ---------------------------------------------------------------------------
CP1252_SPECIAL = {
    0x20AC: 0x80, 0x201A: 0x82, 0x0192: 0x83, 0x201E: 0x84, 0x2026: 0x85, 0x2020: 0x86,
    0x2021: 0x87, 0x02C6: 0x88, 0x2030: 0x89, 0x0160: 0x8A, 0x2039: 0x8B, 0x0152: 0x8C,
    0x017D: 0x8E, 0x2018: 0x91, 0x2019: 0x92, 0x201C: 0x93, 0x201D: 0x94, 0x2022: 0x95,
    0x2013: 0x96, 0x2014: 0x97, 0x02DC: 0x98, 0x2122: 0x99, 0x0161: 0x9A, 0x203A: 0x9B,
    0x0153: 0x9C, 0x017E: 0x9E, 0x0178: 0x9F,
    0x2010: 0xAD,
}

NON_CONNECTING_LETTERS = {'ا', 'أ', 'إ', 'آ', 'د', 'ذ', 'ر', 'ز', 'و', 'ؤ', 'ة', 'ء', 'ى'}

def decode_mojibake(s):
    """
    Decodes Windows-1252 / UTF-8 Mojibake into clean Unicode text.
    Handles contextual undefined character pairs:
    - Ø (0xD8) followed by \u200B / \u2010 -> ح (0xD8 0xAD)
    - Ù (0xD9) followed by \u200B -> ف (0xD9 0x81)
    """
    if not s:
        return ''
    out_bytes = bytearray()
    i = 0
    while i < len(s):
        c = s[i]
        code = ord(c)
        if c == 'Ø' and i + 1 < len(s) and s[i+1] in ('\u200b', '\u2010', '‐'):
            out_bytes.extend([0xd8, 0xad]) # ح
            i += 2
            continue
        if c == 'Ù' and i + 1 < len(s) and s[i+1] == '\u200b':
            out_bytes.extend([0xd9, 0x81]) # ف
            i += 2
            continue
        if c in ('‐', '\u2010'):
            i += 1
            continue
        if code <= 0xFF:
            out_bytes.append(code)
        elif code in CP1252_SPECIAL:
            out_bytes.append(CP1252_SPECIAL[code])
        elif c == '\u200b':
            pass
        else:
            out_bytes.append(0x3F)
        i += 1
    return out_bytes.decode('utf-8', errors='replace')

def reconstruct_arabic_name(raw_name):
    """
    Reconstructs complete Arabic student names by:
    1. Decoding character mojibake.
    2. Joining fragments broken by narrow column CSS character wrapping.
    3. Restoring compound names (عبدالمنعم, أبوالمعاطي, etc.).
    """
    if not raw_name:
        return ''
    lines = [decode_mojibake(line).strip() for line in str(raw_name).splitlines() if line.strip()]
    if not lines:
        return ''
    
    full_tokens = []
    for line in lines:
        tokens = line.split()
        if not full_tokens:
            full_tokens = tokens
        else:
            prev = full_tokens[-1]
            first_curr = tokens[0]
            
            # Check if previous token should be glued to current token:
            glue = False
            if prev in ('م', 'ال', 'الم', 'عبدالم', 'أبوالم', 'محم', 'مح', 'سم'):
                glue = True
            elif prev.endswith(('الم', 'عبدالم', 'أبوالم')):
                glue = True
            elif prev[-1] not in NON_CONNECTING_LETTERS and len(prev) <= 2:
                glue = True
            elif prev[-1] not in NON_CONNECTING_LETTERS and len(first_curr) <= 2:
                glue = True
            elif prev in ('عبد', 'أبو', 'ابن') and first_curr in ('الم', 'ال'):
                glue = True

            if glue:
                full_tokens[-1] = prev + first_curr
                full_tokens.extend(tokens[1:])
            else:
                full_tokens.extend(tokens)
                
    result = ' '.join(full_tokens)
    
    # Post-clean single floating letters that clearly belong to the next word
    words = result.split()
    fixed_words = []
    idx = 0
    while idx < len(words):
        w = words[idx]
        if w in ('م', 'ال', 'الم') and idx + 1 < len(words):
            fixed_words.append(w + words[idx+1])
            idx += 2
        else:
            fixed_words.append(w)
            idx += 1
    return ' '.join(fixed_words)

def clean_cell_text(cell_str):
    """Cleans general table cells, decoding mojibake and normalizing whitespace."""
    if cell_str is None:
        return ''
    s = str(cell_str).strip()
    if not s:
        return ''
    s = s.replace('\r', '').replace('\n', ' ')
    decoded = decode_mojibake(s)
    return ' '.join(decoded.split())

# ---------------------------------------------------------------------------
# PDF Parsing & Data Extraction
# ---------------------------------------------------------------------------
def extract_mark_list_data(pdf_input):
    """
    Parses a mark list PDF from file path or bytes.
    """
    if isinstance(pdf_input, (bytes, bytearray)):
        doc = fitz.open(stream=pdf_input, filetype='pdf')
    else:
        doc = fitz.open(pdf_input)
        
    page = doc[0]
    full_text = page.get_text()

    # Metadata extraction
    course_m = re.search(r'Course:\s*([^\n]+)', full_text)
    subj_m = re.search(r'Subject:\s*([^\n]+)', full_text)
    students_m = re.search(r'Students:\s*([0-9]+)', full_text)
    date_m = re.search(r'Date:\s*([0-9/]+)', full_text)
    deg_m = re.search(r'Degree:\s*([0-9]+)', full_text)
    gen_m = re.search(r'Generated on:\s*([^\n]+)', full_text)

    course_raw = course_m.group(1).strip() if course_m else ''
    course = decode_mojibake(course_raw)
    subject = subj_m.group(1).strip() if subj_m else ''
    date = date_m.group(1).strip() if date_m else ''
    total_degree = deg_m.group(1).strip() if deg_m else '100'
    generated_on = gen_m.group(1).strip() if gen_m else ''

    tabs = page.find_tables()
    if not tabs.tables:
        raise ValueError("Could not find tabular student data in the uploaded PDF.")
    
    extracted_table = tabs.tables[0].extract()
    if not extracted_table or len(extracted_table) < 2:
        raise ValueError("The extracted table has no data rows.")

    raw_headers = extracted_table[0]
    headers = [clean_cell_text(h) for h in raw_headers]

    # Find index for Student Name
    name_col_idx = 2
    for i, h in enumerate(headers):
        if 'name' in h.lower():
            name_col_idx = i
            break

    # Process data rows
    rows = []
    totals = []
    pass_count = 0
    fail_grades = {'ر|F', 'غ|Abs', 'من|W', 'رن|Fr', 'F', 'Abs', 'W', 'Fr', 'F (ر)', 'Abs (غ)', 'W (من)', 'Fr (رن)'}

    for row_idx, raw_row in enumerate(extracted_table[1:], start=1):
        cleaned_row = []
        for col_idx, cell in enumerate(raw_row):
            if col_idx == name_col_idx:
                cleaned_name = reconstruct_arabic_name(cell)
                cleaned_row.append(cleaned_name)
            else:
                cleaned_val = clean_cell_text(cell)
                cleaned_row.append(cleaned_val)
        rows.append(cleaned_row)

        # Total score
        for val in reversed(cleaned_row):
            try:
                num = float(val)
                totals.append(num)
                break
            except ValueError:
                continue

        row_str = ' '.join(cleaned_row)
        is_fail = any(fg in row_str for fg in fail_grades)
        if not is_fail:
            pass_count += 1

    count = len(rows)
    avg_total = round(sum(totals) / len(totals), 1) if totals else 0
    max_total = max(totals) if totals else 0
    min_total = min(totals) if totals else 0
    pass_rate = round((pass_count / count) * 100, 1) if count > 0 else 0

    return {
        "metadata": {
            "course": course,
            "subject": subject,
            "students": count,
            "date": date,
            "degree": total_degree,
            "generated": generated_on
        },
        "headers": headers,
        "rows": rows,
        "summary": {
            "count": count,
            "avg_total": avg_total,
            "max_total": max_total,
            "min_total": min_total,
            "pass_count": pass_count,
            "pass_rate": pass_rate
        }
    }

# ---------------------------------------------------------------------------
# Redesigned Vector PDF Generator
# ---------------------------------------------------------------------------
def generate_redesigned_pdf(data, output_path):
    """
    Generates an executive, beautifully formatted landscape PDF for the Mark List.
    """
    meta = data["metadata"]
    headers = data["headers"]
    rows = data["rows"]
    summary = data["summary"]

    doc = SimpleDocTemplate(
        output_path,
        pagesize=landscape(A4),
        leftMargin=28,
        rightMargin=28,
        topMargin=28,
        bottomMargin=28
    )

    page_width, page_height = landscape(A4)
    content_width = page_width - 56

    styles = getSampleStyleSheet()

    title_style = ParagraphStyle(
        'MainTitle',
        fontName=ARABIC_FONT_NAME,
        fontSize=18,
        leading=22,
        alignment=1,
        textColor=colors.HexColor('#0f172a'),
        fontBold=True
    )

    subtitle_style = ParagraphStyle(
        'Subtitle',
        fontName=ARABIC_FONT_NAME,
        fontSize=11,
        leading=14,
        alignment=1,
        textColor=colors.HexColor('#475569')
    )

    cell_hdr_style = ParagraphStyle(
        'HdrCell',
        fontName=ARABIC_FONT_NAME,
        fontSize=8.5,
        leading=11,
        alignment=1,
        textColor=colors.white,
        fontBold=True
    )

    cell_data_style = ParagraphStyle(
        'DataCell',
        fontName=ARABIC_FONT_NAME,
        fontSize=8.5,
        leading=11,
        alignment=1,
        textColor=colors.HexColor('#1e293b')
    )

    cell_name_style = ParagraphStyle(
        'NameCell',
        fontName=ARABIC_FONT_NAME,
        fontSize=9.5,
        leading=12,
        alignment=2, # Right alignment for Arabic name
        textColor=colors.HexColor('#0f172a')
    )

    card_label_style = ParagraphStyle(
        'CardLabel',
        fontName=ARABIC_FONT_NAME,
        fontSize=8,
        leading=10,
        alignment=1,
        textColor=colors.HexColor('#64748b')
    )

    card_val_style = ParagraphStyle(
        'CardVal',
        fontName=ARABIC_FONT_NAME,
        fontSize=10,
        leading=13,
        alignment=1,
        textColor=colors.HexColor('#0f172a'),
        fontBold=True
    )

    story = []

    # Title & Subtitle
    story.append(Paragraph(ar("Subject Mark List • كشف درجات المقرر"), title_style))
    story.append(Spacer(1, 4))
    story.append(Paragraph(f"<b>{meta['subject']}</b>  —  {ar(meta['course'])}", subtitle_style))
    story.append(Spacer(1, 10))

    # Metadata Cards (Two-tier cell: Label on top, Value below)
    meta_cells = [
        [
            Paragraph("COURSE", card_label_style),
            Paragraph("SUBJECT", card_label_style),
            Paragraph("TOTAL STUDENTS", card_label_style),
            Paragraph("EXAM DATE", card_label_style),
            Paragraph("MAX DEGREE", card_label_style),
        ],
        [
            Paragraph(ar(meta['course']) or "-", card_val_style),
            Paragraph(meta['subject'] or "-", card_val_style),
            Paragraph(str(meta['students']), card_val_style),
            Paragraph(meta['date'] or "-", card_val_style),
            Paragraph(str(meta['degree']), card_val_style),
        ]
    ]
    meta_table = Table(meta_cells, colWidths=[content_width / 5] * 5)
    meta_table.setStyle(TableStyle([
        ('BACKGROUND', (0,0), (-1,-1), colors.HexColor('#f8fafc')),
        ('BOX', (0,0), (-1,-1), 1, colors.HexColor('#e2e8f0')),
        ('ALIGN', (0,0), (-1,-1), 'CENTER'),
        ('VALIGN', (0,0), (-1,-1), 'MIDDLE'),
        ('TOPPADDING', (0,0), (-1,0), 6),
        ('BOTTOMPADDING', (0,0), (-1,0), 2),
        ('TOPPADDING', (0,1), (-1,1), 2),
        ('BOTTOMPADDING', (0,1), (-1,1), 6),
    ]))
    story.append(meta_table)
    story.append(Spacer(1, 12))

    # Column widths calculation
    num_cols = len(headers)
    name_idx = 2
    for i, h in enumerate(headers):
        if 'name' in h.lower():
            name_idx = i
            break

    # Targeted widths:
    # # -> 22, Code -> 62, Name -> 195
    col_widths = [None] * num_cols
    col_widths[0] = 22
    col_widths[1] = 62
    col_widths[name_idx] = 195

    # Check for wide headers like Coursework
    for i, h in enumerate(headers):
        if col_widths[i] is None and 'coursework' in h.lower():
            col_widths[i] = 56

    allocated = sum(w for w in col_widths if w is not None)
    rem_cols = sum(1 for w in col_widths if w is None)
    remaining_w = content_width - allocated
    base_w = remaining_w / max(1, rem_cols)

    for i in range(num_cols):
        if col_widths[i] is None:
            col_widths[i] = base_w

    # Table Header Row
    table_rows = []
    hdr_cells = []
    for h in headers:
        hdr_cells.append(Paragraph(ar(h), cell_hdr_style))
    table_rows.append(hdr_cells)

    # Table Data Rows
    for row in rows:
        row_cells = []
        for col_idx, val in enumerate(row):
            if col_idx == name_idx:
                row_cells.append(Paragraph(ar(val), cell_name_style))
            elif 'grade' in headers[col_idx].lower() and '|' in str(val):
                formatted_g = format_grade(val)
                row_cells.append(Paragraph(formatted_g, cell_data_style))
            else:
                row_cells.append(Paragraph(ar(val), cell_data_style))
        table_rows.append(row_cells)

    marks_table = Table(table_rows, colWidths=col_widths, repeatRows=1)
    
    t_style = [
        ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor('#1e3a8a')),
        ('TEXTCOLOR', (0, 0), (-1, 0), colors.white),
        ('ALIGN', (0, 0), (-1, 0), 'CENTER'),
        ('VALIGN', (0, 0), (-1, -1), 'MIDDLE'),
        ('TOPPADDING', (0, 0), (-1, -1), 4.5),
        ('BOTTOMPADDING', (0, 0), (-1, -1), 4.5),
        ('LEFTPADDING', (0, 0), (-1, -1), 4),
        ('RIGHTPADDING', (0, 0), (-1, -1), 4),
        ('GRID', (0, 0), (-1, -1), 0.5, colors.HexColor('#cbd5e1')),
    ]

    for r in range(1, len(table_rows)):
        bg_col = colors.HexColor('#ffffff') if r % 2 != 0 else colors.HexColor('#f8fafc')
        t_style.append(('BACKGROUND', (0, r), (-1, r), bg_col))

    marks_table.setStyle(TableStyle(t_style))
    story.append(marks_table)
    story.append(Spacer(1, 12))

    # Summary KPI Statistics Footer Block
    kpi_labels = [
        Paragraph("CLASS AVERAGE", card_label_style),
        Paragraph("HIGHEST MARK", card_label_style),
        Paragraph("LOWEST MARK", card_label_style),
        Paragraph("STUDENTS PASSED", card_label_style),
        Paragraph("PASSING RATE", card_label_style),
    ]
    kpi_vals = [
        Paragraph(f"<b>{summary['avg_total']}</b>", card_val_style),
        Paragraph(f"<b>{summary['max_total']}</b>", card_val_style),
        Paragraph(f"<b>{summary['min_total']}</b>", card_val_style),
        Paragraph(f"<b>{summary['pass_count']} / {summary['count']}</b>", card_val_style),
        Paragraph(f"<b>{summary['pass_rate']}%</b>", card_val_style),
    ]
    summary_table = Table([kpi_labels, kpi_vals], colWidths=[content_width / 5] * 5)
    summary_table.setStyle(TableStyle([
        ('BACKGROUND', (0,0), (-1,-1), colors.HexColor('#e0f2fe')),
        ('BOX', (0,0), (-1,-1), 1, colors.HexColor('#7dd3fc')),
        ('ALIGN', (0,0), (-1,-1), 'CENTER'),
        ('VALIGN', (0,0), (-1,-1), 'MIDDLE'),
        ('TOPPADDING', (0,0), (-1,0), 5),
        ('BOTTOMPADDING', (0,0), (-1,0), 2),
        ('TOPPADDING', (0,1), (-1,1), 2),
        ('BOTTOMPADDING', (0,1), (-1,1), 5),
    ]))
    story.append(summary_table)
    story.append(Spacer(1, 8))

    footer_text = f"Report Re-engineered by PDF Arabic Fixer • Source Generated: {meta['generated'] or meta['date']}"
    story.append(Paragraph(footer_text, subtitle_style))

    doc.build(story)
    return output_path

# ---------------------------------------------------------------------------
# Excel Workbook Generator
# ---------------------------------------------------------------------------
def export_to_excel(data, output_path):
    """
    Exports the restored student marks to an aesthetically formatted Excel workbook.
    """
    meta = data["metadata"]
    headers = data["headers"]
    rows = data["rows"]
    summary = data["summary"]

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Mark List"
    ws.views.sheetView[0].showGridLines = True

    title_font = Font(name='Arial', size=14, bold=True, color='1E3A8A')
    subtitle_font = Font(name='Arial', size=11, italic=True, color='475569')
    hdr_font = Font(name='Arial', size=10, bold=True, color='FFFFFF')
    hdr_fill = PatternFill(start_color='1E3A8A', end_color='1E3A8A', fill_type='solid')
    
    border_side = Side(border_style='thin', color='CBD5E1')
    cell_border = Border(left=border_side, right=border_side, top=border_side, bottom=border_side)
    zebra_fill = PatternFill(start_color='F8FAFC', end_color='F8FAFC', fill_type='solid')

    # Row 1: Title
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=len(headers))
    cell = ws.cell(row=1, column=1, value="Subject Mark List - كشف درجات المقرر")
    cell.font = title_font
    cell.alignment = Alignment(horizontal='center', vertical='center')

    # Row 2: Subtitle
    ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=len(headers))
    cell = ws.cell(row=2, column=1, value=f"{meta['subject']} | {meta['course']} (Exam Date: {meta['date']})")
    cell.font = subtitle_font
    cell.alignment = Alignment(horizontal='center', vertical='center')

    ws.append([]) # Row 3 blank

    # Row 4: Table Headers
    ws.append(headers)
    hdr_row = 4
    for col_num in range(1, len(headers) + 1):
        c = ws.cell(row=hdr_row, column=col_num)
        c.font = hdr_font
        c.fill = hdr_fill
        c.alignment = Alignment(horizontal='center', vertical='center')
        c.border = cell_border

    # Data Rows
    current_row = 5
    for row in rows:
        ws.append(row)
        for col_num, val in enumerate(row, start=1):
            c = ws.cell(row=current_row, column=col_num)
            c.border = cell_border
            if col_num == 3:
                c.alignment = Alignment(horizontal='right', vertical='center')
            else:
                c.alignment = Alignment(horizontal='center', vertical='center')
            if (current_row - 4) % 2 == 0:
                c.fill = zebra_fill
        current_row += 1

    current_row += 1

    # Summary Row
    summary_txt = f"Average: {summary['avg_total']} | Highest: {summary['max_total']} | Lowest: {summary['min_total']} | Passed: {summary['pass_count']}/{summary['count']} ({summary['pass_rate']}%)"
    ws.merge_cells(start_row=current_row, start_column=1, end_row=current_row, end_column=len(headers))
    sum_cell = ws.cell(row=current_row, column=1, value=summary_txt)
    sum_cell.font = Font(name='Arial', size=10, bold=True, color='0369A1')
    sum_cell.fill = PatternFill(start_color='E0F2FE', end_color='E0F2FE', fill_type='solid')
    sum_cell.alignment = Alignment(horizontal='center', vertical='center')

    # Auto-adjust column widths
    for col in ws.columns:
        max_len = 0
        col_letter = get_column_letter(col[0].column)
        for c in col:
            val_str = str(c.value or '')
            max_len = max(max_len, len(val_str))
        ws.column_dimensions[col_letter].width = max(max_len + 4, 12)
    ws.column_dimensions['C'].width = 38

    wb.save(output_path)
    return output_path
